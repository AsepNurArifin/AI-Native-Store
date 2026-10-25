#!/usr/bin/env python3
"""
build-dokumen.py — Konverter Markdown -> DOCX + PDF untuk dokumen tugas Shaf Store.

Pemakaian:
    python scripts/build-dokumen.py

Hasil:
    docs/export/Dokumen-Shaf-Store.pdf      +  .docx
    docs/export/Naskah-Konten.pdf           +  .docx

Dependensi:  python -m pip install python-docx xhtml2pdf
"""

from __future__ import annotations

import html
import pathlib
import re
import sys

# ---------------------------------------------------------------- konfigurasi

ROOT = pathlib.Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
EXPORT = DOCS / "export"

SUMBER = [
    DOCS / "Dokumen-Shaf-Store.md",
    DOCS / "Naskah-Konten.md",
]

# ---------------------------------------------------------------- parser

INLINE_RE = re.compile(
    r"(\*\*.+?\*\*"          # **bold**
    r"|__.+?__"              # __bold__
    r"|`[^`]+`"              # `code`
    r"|(?<!\*)\*[^*]+?\*(?!\*)"   # *italic*
    r"|(?<!_)_[^_]+?_(?!_)"       # _italic_
    r"|\[[^\]]+\]\([^)]+\)"       # [teks](url)
    r")"
)


def parse_blocks(text: str) -> list[dict]:
    """Pecah markdown menjadi blok-blok terstruktur."""
    lines = text.replace("\r\n", "\n").split("\n")
    blocks: list[dict] = []
    i = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # kosong
        if stripped == "":
            i += 1
            continue

        # horizontal rule
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            blocks.append({"type": "hr"})
            i += 1
            continue

        # code fence
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1  # tutup fence
            blocks.append({"type": "code", "text": "\n".join(buf)})
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            blocks.append({"type": "h", "level": len(m.group(1)), "text": m.group(2).strip()})
            i += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            blocks.append({"type": "quote", "text": "\n".join(buf).strip()})
            continue

        # tabel: baris pertama memakai '|' dan baris kedua adalah pemisah
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(
            r"^\s*\|[\s:|-]+\|?\s*$", lines[i + 1]
        ):
            header = _split_row(stripped)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_row(lines[i].strip()))
                i += 1
            blocks.append({"type": "table", "header": header, "rows": rows})
            continue

        # list
        if re.match(r"^\s*[-*+]\s+", line):
            items = []
            while i < len(lines) and re.match(r"^\s*[-*+]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*+]\s+", "", lines[i]))
                i += 1
            blocks.append({"type": "ul", "items": items})
            continue

        if re.match(r"^\s*\d+[.)]\s+", line):
            items = []
            while i < len(lines) and re.match(r"^\s*\d+[.)]\s+", lines[i]):
                items.append(re.sub(r"^\s*\d+[.)]\s+", "", lines[i]))
                i += 1
            blocks.append({"type": "ol", "items": items})
            continue

        # paragraf (gabung baris berurutan)
        buf = []
        while i < len(lines):
            cur = lines[i]
            s = cur.strip()
            if (
                s == ""
                or s.startswith("#")
                or s.startswith(">")
                or s.startswith("```")
                or s.startswith("|")
                or re.match(r"^\s*[-*+]\s+", cur)
                or re.match(r"^\s*\d+[.)]\s+", cur)
                or re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", s)
            ):
                break
            buf.append(s)
            i += 1
        if buf:
            blocks.append({"type": "p", "text": " ".join(buf)})

    return blocks


def _split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


# ---------------------------------------------------------------- inline

def parse_inline(text: str) -> list[tuple[str, str]]:
    """Kembalikan daftar (kind, value): 'text' | 'bold' | 'italic' | 'code' | 'link'."""
    out: list[tuple[str, str]] = []
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            out.append(("bold", part[2:-2]))
        elif part.startswith("__") and part.endswith("__") and len(part) > 4:
            out.append(("bold", part[2:-2]))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            out.append(("code", part[1:-1]))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append(("italic", part[1:-1]))
        elif part.startswith("_") and part.endswith("_") and len(part) > 2:
            out.append(("italic", part[1:-1]))
        else:
            m = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
            if m:
                out.append(("link", f"{m.group(1)}\x00{m.group(2)}"))
            else:
                out.append(("text", part))
    return out


# ---------------------------------------------------------------- DOCX

def build_docx(blocks: list[dict], out_path: pathlib.Path, title: str) -> None:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor

    doc = Document()

    # A4 + margin
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)

    # dasar font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    seen_h1 = False

    for block in blocks:
        t = block["type"]

        if t == "hr":
            p = doc.add_paragraph()
            p.add_run("\u2500" * 46).font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
            continue

        if t == "h":
            lvl = block["level"]
            if lvl == 1 and not seen_h1:
                para = doc.add_paragraph()
                run = para.add_run(_plain(block["text"]))
                run.bold = True
                run.font.size = Pt(22)
                run.font.color.rgb = RGBColor(0x1F, 0x2A, 0x37)
                seen_h1 = True
            else:
                doc.add_heading(_plain(block["text"]), level=min(lvl, 3))
            continue

        if t == "p":
            p = doc.add_paragraph()
            _add_runs(p, block["text"], Pt(11))
            p.paragraph_format.space_after = Pt(8)
            continue

        if t == "quote":
            for line in block["text"].split("\n"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                _add_runs(p, line, Pt(10.5), italic=True)
                for r in p.runs:
                    r.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
            continue

        if t == "code":
            for line in block["text"].split("\n"):
                p = doc.add_paragraph()
                run = p.add_run(line if line else " ")
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.space_after = Pt(0)
            continue

        if t in ("ul", "ol"):
            style_name = "List Bullet" if t == "ul" else "List Number"
            for item in block["items"]:
                try:
                    p = doc.add_paragraph(style=style_name)
                except KeyError:
                    p = doc.add_paragraph(style="List Bullet")
                _add_runs(p, item, Pt(11))
                p.paragraph_format.space_after = Pt(3)
            continue

        if t == "table":
            header = block["header"]
            rows = block["rows"]
            cols = max([len(header)] + [len(r) for r in rows]) if rows else len(header)
            cols = max(cols, 1)
            table = doc.add_table(rows=1, cols=cols)
            try:
                table.style = "Table Grid"
            except KeyError:
                pass

            for idx, cell_text in enumerate(header[:cols]):
                cell = table.rows[0].cells[idx]
                cell.text = ""
                _add_runs(cell.paragraphs[0], cell_text, Pt(10), bold=True)

            for row in rows:
                cells = table.add_row().cells
                for idx in range(cols):
                    val = row[idx] if idx < len(row) else ""
                    cells[idx].text = ""
                    _add_runs(cells[idx].paragraphs[0], val, Pt(10))

            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

    doc.save(out_path)


def _plain(text: str) -> str:
    return re.sub(r"[*_`]", "", re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text))


def _add_runs(paragraph, text: str, size, bold: bool = False, italic: bool = False) -> None:
    for kind, value in parse_inline(text):
        if kind == "link":
            label, url = value.split("\x00")
            run = paragraph.add_run(label)
            run.underline = True
            run.font.size = size
            run.bold = bold
            run.italic = italic
            continue

        run = paragraph.add_run(value)
        run.font.size = size
        run.bold = bold or (kind == "bold")
        run.italic = italic or (kind == "italic")
        if kind == "code":
            run.font.name = "Consolas"


# ---------------------------------------------------------------- HTML

# Font bawaan xhtml2pdf (Helvetica/Courier) tidak memiliki glyph untuk simbol
# berikut. Hanya jalur PDF yang ditransliterasi; DOCX memakai UTF-8 penuh.
PDF_CHAR_MAP = {
    "\u2248": "~",      # ≈
    "\u2502": "|",      # │
    "\u2503": "|",      # ┃
    "\u25bc": "v",      # ▼
    "\u25ba": ">",      # ►
    "\u2500": "-",      # ─
    "\u2501": "-",      # ━
    "\u2b50": "*",      # ⭐
    "\u2192": "->",     # →
    "\u2190": "<-",     # ←
    "\u251c": "+",      # ├
    "\u2514": "\\",     # └
    "\u2524": "+",      # ┤
    "\u250c": "+",      # ┌
    "\u2510": "+",      # ┐
    "\u2518": "+",      # ┘
    "\u2022": "-",      # •
}


def _pdf_safe(text: str) -> str:
    for bad, good in PDF_CHAR_MAP.items():
        text = text.replace(bad, good)
    return text


def inline_html(text: str) -> str:
    text = _pdf_safe(text)
    out = []
    for kind, value in parse_inline(text):
        if kind == "bold":
            out.append(f"<b>{html.escape(value)}</b>")
        elif kind == "italic":
            out.append(f"<i>{html.escape(value)}</i>")
        elif kind == "code":
            out.append(f'<font face="Courier">{html.escape(value)}</font>')
        elif kind == "link":
            label, url = value.split("\x00")
            out.append(f'<a href="{html.escape(url)}">{html.escape(label)}</a>')
        else:
            out.append(html.escape(value))
    return "".join(out)


def blocks_to_html(blocks: list[dict]) -> str:
    out: list[str] = []
    for block in blocks:
        t = block["type"]

        if t == "hr":
            out.append('<hr/>')

        elif t == "h":
            lvl = block["level"]
            sizes = {1: "20pt", 2: "15pt", 3: "12.5pt"}
            size = sizes.get(min(lvl, 3), "12pt")
            out.append(
                f'<p style="font-size:{size};font-weight:bold;margin-top:{14 if lvl > 1 else 6}px;'
                f'margin-bottom:5px;">{inline_html(block["text"])}</p>'
            )

        elif t == "p":
            out.append(f'<p style="margin-bottom:7px;line-height:1.5;">{inline_html(block["text"])}</p>')

        elif t == "quote":
            inner = "<br/>".join(inline_html(l) for l in block["text"].split("\n"))
            out.append(
                '<table width="100%" cellpadding="0" cellspacing="0"><tr>'
                '<td width="5" bgcolor="#C9CFD6"></td>'
                '<td width="8"></td>'
                f'<td style="color:#4B5563;font-style:italic;font-size:10pt;">{inner}</td>'
                "</tr></table>"
            )

        elif t == "code":
            body = "<br/>".join(html.escape(_pdf_safe(l)) for l in block["text"].split("\n"))
            out.append(
                f'<p style="font-family:Courier;font-size:8.5pt;margin-left:10px;'
                f'line-height:1.35;">{body}</p>'
            )

        elif t == "ul":
            items = "".join(f"<li>{inline_html(i)}</li>" for i in block["items"])
            out.append(f'<ul style="margin-bottom:7px;line-height:1.45;">{items}</ul>')

        elif t == "ol":
            items = "".join(f"<li>{inline_html(i)}</li>" for i in block["items"])
            out.append(f'<ol style="margin-bottom:7px;line-height:1.45;">{items}</ol>')

        elif t == "table":
            header = block["header"]
            rows = block["rows"]
            cols = max([len(header)] + [len(r) for r in rows]) if rows else len(header)
            cols = max(cols, 1)

            th = "".join(
                f'<th align="left" bgcolor="#E8ECF1" style="padding:5px 7px;'
                f'border:0.5px solid #B9C2CC;font-size:9pt;">{inline_html(header[i] if i < len(header) else "")}</th>'
                for i in range(cols)
            )
            body_rows = []
            for r in rows:
                tds = "".join(
                    f'<td style="padding:5px 7px;border:0.5px solid #B9C2CC;font-size:9pt;">'
                    f'{inline_html(r[i] if i < len(r) else "")}</td>'
                    for i in range(cols)
                )
                body_rows.append(f"<tr>{tds}</tr>")

            out.append(
                '<table width="100%" cellpadding="0" cellspacing="0" '
                'style="margin-bottom:10px;">'
                f"<tr>{th}</tr>{''.join(body_rows)}</table>"
            )

    return "\n".join(out)


def build_pdf(blocks: list[dict], out_path: pathlib.Path) -> None:
    from xhtml2pdf import pisa

    body = blocks_to_html(blocks)
    doc_html = f"""<!DOCTYPE html>
<html>
<head><meta charset="utf-8"/>
<style>
  body {{ font-family: Helvetica, Arial, sans-serif; font-size: 10.5pt; color: #1F2A37; }}
  a {{ color: #1F2A37; text-decoration: underline; }}
</style>
</head>
<body>{body}</body>
</html>"""

    with open(out_path, "wb") as fh:
        status = pisa.CreatePDF(doc_html, dest=fh, encoding="utf-8")

    if status.err:
        raise RuntimeError(f"Gagal membuat PDF: {out_path.name}")


# ---------------------------------------------------------------- main

def main() -> int:
    EXPORT.mkdir(parents=True, exist_ok=True)

    for md_path in SUMBER:
        if not md_path.exists():
            print(f"  lewati (tidak ada): {md_path.name}")
            continue

        text = md_path.read_text(encoding="utf-8")
        blocks = parse_blocks(text)
        stem = md_path.stem

        docx_path = EXPORT / f"{stem}.docx"
        pdf_path = EXPORT / f"{stem}.pdf"

        build_docx(blocks, docx_path, stem)
        build_pdf(blocks, pdf_path)

        print(f"  {md_path.name}")
        print(f"      -> {docx_path.relative_to(ROOT)}  ({docx_path.stat().st_size:,} bytes)")
        print(f"      -> {pdf_path.relative_to(ROOT)}   ({pdf_path.stat().st_size:,} bytes)")

    print(f"\nSelesai. Output di: {EXPORT.relative_to(ROOT)}/")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:  # noqa: BLE001
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
